"""
load.py - Load microbiome rollup xlsx files into Supabase Postgres.

Usage:
    python load.py path/to/saccharomyces_boulardii_rollup.xlsx
    python load.py --all data/                  # load every *_rollup.xlsx in a folder
    python load.py --dry-run path/to/file.xlsx  # parse only, no DB writes

Reads SUPABASE_DB_URL from .env in the same directory as this script.
Re-loading the same file is safe: it deletes existing rows for that
modifier (cascade) and re-inserts.
"""

import argparse
import os
import sys
from pathlib import Path

import openpyxl
import psycopg2
from psycopg2.extras import execute_values
from dotenv import load_dotenv


def parse_db_url(url):
    """
    Parse postgresql://user:password@host:port/db into psycopg2 kwargs.
    Done manually so passwords containing %, +, :, etc. work without URL-encoding.
    """
    if not url.startswith(("postgres://", "postgresql://")):
        raise ValueError("DB URL must start with postgres:// or postgresql://")
    _, rest = url.split("://", 1)
    if "@" not in rest:
        raise ValueError("DB URL missing host part")
    userinfo, hostpart = rest.rsplit("@", 1)
    if ":" in userinfo:
        username, password = userinfo.split(":", 1)
    else:
        username, password = userinfo, None
    if "/" in hostpart:
        hostport, db = hostpart.split("/", 1)
        if "?" in db:
            db = db.split("?", 1)[0]
    else:
        hostport, db = hostpart, "postgres"
    if ":" in hostport:
        host, port_str = hostport.rsplit(":", 1)
        port = int(port_str)
    else:
        host, port = hostport, 5432
    return {"host": host, "port": port, "dbname": db, "user": username, "password": password}


SHEET_TO_TABLE = {
    "Rollup": "rollup",
    "Composite Breakdown": "composite_breakdown",
    "Per-Study Findings": "per_study_findings",
    "Studies": "studies",
    "Excluded Studies": "excluded_studies",
    "Per-Study Narrative": "per_study_narrative",
    "Per-Finding Narrative": "per_finding_narrative",
    "Validation Issues": "validation_issues",
}

INGESTION_ORDER = [
    "studies", "excluded_studies", "rollup", "composite_breakdown",
    "per_study_findings", "per_study_narrative", "per_finding_narrative",
    "validation_issues",
]

EXPECTED_COLS = {
    "studies": [
        "modifier", "study_id", "cohort_cluster_id", "first_author", "year",
        "journal", "doi", "study_type", "quality_tier", "population",
        "sample_size", "microbiome_method", "sample_type", "intervention",
        "duration", "is_disease_state_model", "disease_model",
        "include_for_scoring", "n_findings_rows", "alpha_diversity_change",
        "beta_diversity_change",
    ],
    "excluded_studies": [
        "modifier", "study_id", "first_author", "year",
        "exclusion_reason", "evidence_basis",
    ],
    "rollup": [
        "modifier", "taxon_listed_name", "taxon_id", "direction",
        "increase_effect", "increase_confidence", "decrease_effect",
        "decrease_confidence", "evidence_tier", "n_studies_contributing",
        "n_human_direct", "n_animal_direct", "n_disease_state_model",
        "discordance_warning", "cohort_cluster_summary",
        "contributing_study_ids", "evidence_basis",
    ],
    "composite_breakdown": [
        "modifier", "composite_taxon", "member_taxon", "member_taxon_id",
        "n_studies_total", "n_studies_increase", "n_studies_decrease",
        "n_studies_mixed", "n_studies_insufficient", "increase_effect",
        "increase_confidence", "decrease_effect", "decrease_confidence",
        "contributing_study_ids", "sample_quotes",
    ],
    "per_study_findings": [
        "modifier", "study_id", "cohort_cluster_id", "study_first_author",
        "study_year", "study_type", "sample_size", "is_disease_state_model",
        "disease_model", "taxon_listed_name", "taxon_id", "intervention_arm",
        "direction", "effect_magnitude", "fold_change", "p_value",
        "statistical_significance", "multiple_comparison_correction",
        "comparison_type", "low_signal", "captured_via",
        "computed_increase_effect_size", "computed_increase_confidence",
        "computed_decrease_effect_size", "computed_decrease_confidence",
        "computed_evidence_tier", "computed_quality_tier", "baseline_value",
        "post_intervention_value", "timepoint", "subgroup_context",
        "mechanism_notes", "paper_interpretation", "additional_quotes",
        "direction_quote", "quote_source", "evidence_basis", "notes",
    ],
    "per_study_narrative": [
        "modifier", "study_id", "first_author", "year", "study_type",
        "paper_discussion_summary", "off_canonical_taxa_observed",
        "raw_extraction_notes", "study_design_details", "statistical_methods",
        "co_interventions", "dose_details", "limitations",
        "additional_diversity_notes", "notes",
    ],
    "per_finding_narrative": [
        "modifier", "study_id", "taxon_listed_name", "intervention_arm",
        "direction", "mechanism_notes", "paper_interpretation",
        "baseline_value", "post_intervention_value", "timepoint",
        "subgroup_context", "evidence_basis", "additional_quotes", "notes",
    ],
    "validation_issues": [
        "modifier", "file", "error_type", "error_detail",
    ],
}

INT_COLS = {
    "studies": {"year", "sample_size", "n_findings_rows"},
    "excluded_studies": {"year"},
    "rollup": {"n_studies_contributing", "n_human_direct", "n_animal_direct", "n_disease_state_model"},
    "composite_breakdown": {"n_studies_total", "n_studies_increase", "n_studies_decrease", "n_studies_mixed", "n_studies_insufficient"},
    "per_study_findings": {"study_year", "sample_size"},
    "per_study_narrative": {"year"},
    "per_finding_narrative": set(),
    "validation_issues": set(),
}

SMALLINT_COLS = {
    "rollup": {"increase_effect", "increase_confidence", "decrease_effect", "decrease_confidence"},
    "composite_breakdown": {"increase_effect", "increase_confidence", "decrease_effect", "decrease_confidence"},
    "per_study_findings": {"computed_increase_effect_size", "computed_increase_confidence", "computed_decrease_effect_size", "computed_decrease_confidence"},
}

BOOL_COLS = {
    "studies": {"is_disease_state_model"},
    "per_study_findings": {"is_disease_state_model"},
}


def parse_filename(path):
    stem = path.stem
    if stem.lower().endswith("_rollup"):
        stem = stem[: -len("_rollup")]
    display_raw = stem.replace("_", " ")
    display_name = display_raw[:1].upper() + display_raw[1:] if display_raw else display_raw
    modifier = stem.lower().replace(" ", "_")
    return modifier, display_name


def coerce(value, kind):
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    if kind in ("int", "smallint"):
        try:
            return int(value)
        except (ValueError, TypeError):
            try:
                return int(float(value))
            except (ValueError, TypeError):
                return None
    if kind == "bool":
        if isinstance(value, bool):
            return value
        s = str(value).strip().lower()
        if s in ("true", "1", "yes", "y", "t"):
            return True
        if s in ("false", "0", "no", "n", "f"):
            return False
        return None
    return str(value)


def column_kind(table, col):
    if col in BOOL_COLS.get(table, set()):
        return "bool"
    if col in INT_COLS.get(table, set()):
        return "int"
    if col in SMALLINT_COLS.get(table, set()):
        return "smallint"
    return "text"


def read_sheet_rows(ws, table, modifier):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).lower() if h is not None else "" for h in rows[0]]
    expected_set = set(EXPECTED_COLS[table])
    out = []
    for raw in rows[1:]:
        if all(c is None or (isinstance(c, str) and c.strip() == "") for c in raw):
            continue
        record = {"modifier": modifier}
        for h, v in zip(headers, raw):
            if not h or h not in expected_set:
                continue
            record[h] = coerce(v, column_kind(table, h))
        out.append(record)
    return out


def insert_rows(conn, table, rows):
    if not rows:
        return
    cols = EXPECTED_COLS[table]
    sql = f"INSERT INTO {table} ({', '.join(cols)}) VALUES %s"
    values = [tuple(r.get(c) for c in cols) for r in rows]
    with conn.cursor() as cur:
        execute_values(cur, sql, values, page_size=500)


def load_file(path, conn, dry_run=False):
    print(f"\n=== {path.name} ===")
    if not path.exists():
        print(f"  ERROR: file not found")
        return

    modifier, display_name = parse_filename(path)
    print(f"  modifier:     {modifier}")
    print(f"  display_name: {display_name}")

    wb = openpyxl.load_workbook(path, data_only=True)
    parsed = {}
    for sheet_name, table in SHEET_TO_TABLE.items():
        if sheet_name not in wb.sheetnames:
            print(f"  WARN: sheet '{sheet_name}' not in workbook, skipping")
            parsed[table] = []
            continue
        rows = read_sheet_rows(wb[sheet_name], table, modifier)
        parsed[table] = rows
        print(f"  {sheet_name:<25} -> {table:<25} {len(rows):>4} rows")

    if dry_run:
        print("  (dry run, no DB writes)")
        return

    print(f"  deleting existing rows for modifier='{modifier}' (cascade)...")
    with conn.cursor() as cur:
        cur.execute("DELETE FROM modifiers WHERE modifier = %s", (modifier,))
        cur.execute(
            "INSERT INTO modifiers (modifier, display_name, source_file) VALUES (%s, %s, %s)",
            (modifier, display_name, path.name),
        )

    for table in INGESTION_ORDER:
        rows = parsed.get(table, [])
        if not rows:
            continue
        insert_rows(conn, table, rows)
        print(f"  inserted {len(rows):>4} into {table}")

    conn.commit()
    print(f"  done.")


def main():
    parser = argparse.ArgumentParser(description="Load rollup xlsx files into Supabase Postgres.")
    parser.add_argument("path", nargs="?", help="Path to xlsx file, or folder if --all is set.")
    parser.add_argument("--all", action="store_true", help="Treat path as a folder.")
    parser.add_argument("--dry-run", action="store_true", help="Parse files but do not write to DB.")
    args = parser.parse_args()

    if not args.path:
        parser.error("provide a path to an xlsx file (or a folder with --all)")

    target = Path(args.path)

    if args.all:
        if not target.is_dir():
            parser.error(f"--all expects a folder, got {target}")
        files = sorted(target.glob("*_rollup.xlsx"))
        if not files:
            parser.error(f"no *_rollup.xlsx files found in {target}")
    else:
        files = [target]

    conn = None
    if not args.dry_run:
        load_dotenv(Path(__file__).parent / ".env")
        db_url = os.environ.get("SUPABASE_DB_URL")
        if not db_url:
            print("ERROR: SUPABASE_DB_URL must be set in .env")
            print("Get it from Supabase: Project Settings > Database > Connection string > Session pooler")
            sys.exit(1)
        try:
            conn = psycopg2.connect(**parse_db_url(db_url))
        except (psycopg2.OperationalError, ValueError) as e:
            print(f"ERROR: could not connect to Supabase Postgres: {e}")
            sys.exit(1)

    try:
        for f in files:
            load_file(f, conn, dry_run=args.dry_run)
    finally:
        if conn is not None:
            conn.close()

    print("\nAll done.")


if __name__ == "__main__":
    main()
